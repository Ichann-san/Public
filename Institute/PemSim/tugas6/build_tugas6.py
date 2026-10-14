from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Tugas 6 - Input Analysis.xlsx"

ACCIDENTS = [
    88.3, 40.7, 36.3, 27.3, 36.8,
    91.7, 67.3, 7.0, 45.2, 23.3,
    98.8, 90.1, 17.2, 23.7, 97.4,
    32.4, 87.8, 69.8, 62.6, 99.7,
    20.6, 73.1, 21.6, 6.0, 45.3,
    76.6, 73.2, 27.3, 87.6, 87.2,
]
TIMES = [
    1.88, .54, 1.90, .15, .02, 2.81, 1.50, .53, 2.62, 2.67,
    3.53, .53, 1.80, .79, .21, .80, .26, .63, .36, 2.03,
    1.42, 1.28, .82, 2.16, .05, .04, 1.49, .66, 2.03, 1.00,
    .39, .34, .01, .10, 1.10, .24, .26, .45, .17, 4.29,
    .80, 5.50, 4.91, .35, .36, .90, 1.03, 1.73, .38, .48,
]

NAVY = "203047"
BLUE = "315D79"
PALE = "EAF1F5"
PALE_ALT = "F6F9FB"
LINE = "D5E0E6"
TEXT = "243442"
INPUT = "1B63A0"
GREEN = "236849"


def base_sheet(ws, title, subtitle, last_col, last_row, orientation="landscape"):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=last_col)
    ws.cell(1, 2, title)
    ws.cell(1, 2).font = Font(name="Arial", size=16, bold=True, color=NAVY)
    ws.cell(1, 2).alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 29
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=last_col)
    ws.cell(2, 2, subtitle)
    ws.cell(2, 2).font = Font(name="Arial", size=9, color=BLUE)
    ws.row_dimensions[2].height = 19
    for row in ws.iter_rows(min_row=3, max_row=last_row, max_col=last_col):
        for c in row:
            c.font = Font(name="Arial", size=9, color=TEXT)
            c.alignment = Alignment(vertical="center")
    ws.column_dimensions["A"].width = 2.8
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = orientation
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = .25
    ws.page_margins.right = .25
    ws.page_margins.top = .35
    ws.page_margins.bottom = .35
    ws.oddFooter.center.text = "Tugas 6 | Input Analysis"
    ws.oddFooter.center.size = 8
    ws.print_area = f"A1:{get_column_letter(last_col)}{last_row}"
    ws.freeze_panes = "B6"


def header(ws, row, start, labels):
    for col, label in enumerate(labels, start):
        c = ws.cell(row, col, label)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = Border(bottom=Side(style="thin", color=NAVY))
    ws.row_dimensions[row].height = 29


def body(ws, start_row, end_row, start_col, end_col):
    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            c = ws.cell(row, col)
            if row % 2 == 0:
                c.fill = PatternFill("solid", fgColor=PALE_ALT)
            c.border = Border(bottom=Side(style="hair", color=LINE))
            c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 18


def note(ws, cell_range, text):
    ws.merge_cells(cell_range)
    c = ws[cell_range.split(":")[0]]
    c.value = text
    c.font = Font(name="Arial", size=8, color=BLUE)
    c.alignment = Alignment(wrap_text=True, vertical="center")


def chart(ws, category_col, value_col, first_row, last_row, anchor, title, width, height):
    c = BarChart()
    c.type = "col"
    c.style = 10
    c.title = title
    c.y_axis.title = "Frekuensi"
    c.x_axis.title = "Interval"
    c.height = height
    c.width = width
    c.legend = None
    c.add_data(Reference(ws, min_col=value_col, min_row=first_row - 1, max_row=last_row), titles_from_data=True)
    c.set_categories(Reference(ws, min_col=category_col, min_row=first_row, max_row=last_row))
    c.dLbls = DataLabelList()
    c.dLbls.showVal = True
    ws.add_chart(c, anchor)


wb = Workbook()
acc = wb.active
acc.title = "1 Accident Data"
ks = wb.create_sheet("1 KS Test")
time = wb.create_sheet("2 Time Data")
chi = wb.create_sheet("2 Chi Square")

# Question 1: raw values, descriptive statistics, and empirical distribution.
base_sheet(acc, "1  Accident locations", "30 injury or fatal accidents on a 100 km road | September", 14, 38)
for col, width in {"B": 8, "C": 13, "D": 3, "E": 19, "F": 13, "G": 3,
                   "H": 16, "I": 12, "J": 10, "K": 10, "L": 10, "M": 10, "N": 8}.items():
    acc.column_dimensions[col].width = width
header(acc, 5, 2, ["No.", "Location (km)"])
body(acc, 6, 35, 2, 3)
for idx, value in enumerate(ACCIDENTS, 1):
    row = idx + 5
    acc.cell(row, 2, idx)
    c = acc.cell(row, 3, value)
    c.number_format = "0.0"
    c.font = Font(name="Arial", size=9, color=INPUT)

header(acc, 5, 5, ["Statistical summary", "Value"])
summary_a = [
    ("Count", "=COUNT(C6:C35)", "0"),
    ("Sum (km)", "=SUM(C6:C35)", "0.0"),
    ("Mean (km)", "=AVERAGE(C6:C35)", "0.000"),
    ("Median (km)", "=MEDIAN(C6:C35)", "0.000"),
    ("Mode (km)", "=MODE(C6:C35)", "0.0"),
    ("Minimum (km)", "=MIN(C6:C35)", "0.0"),
    ("Maximum (km)", "=MAX(C6:C35)", "0.0"),
    ("Range (km)", "=F12-F11", "0.0"),
    ("Q1 (km)", "=QUARTILE(C6:C35,1)", "0.000"),
    ("Q3 (km)", "=QUARTILE(C6:C35,3)", "0.000"),
    ("IQR (km)", "=F15-F14", "0.000"),
    ("Sample variance", "=VAR(C6:C35)", "0.000"),
    ("Sample SD (km)", "=STDEV(C6:C35)", "0.000"),
    ("CV", "=F18/F8", "0.0%"),
    ("Skewness", "=SKEW(C6:C35)", "0.000"),
    ("Excess kurtosis", "=KURT(C6:C35)", "0.000"),
]
body(acc, 6, 21, 5, 6)
for row, (label, formula, fmt) in enumerate(summary_a, 6):
    acc.cell(row, 5, label)
    acc.cell(row, 6, formula).number_format = fmt

header(acc, 5, 8, ["Interval (km)", "Frequency", "Relative", "Cumulative"])
body(acc, 6, 15, 8, 11)
for j in range(10):
    r = 6 + j
    low, high = j * 10, (j + 1) * 10
    acc.cell(r, 8, f"{low}-{high}")
    acc.cell(r, 9, f'=COUNTIFS($C$6:$C$35,">="&{low},$C$6:$C$35,"<"&{high})')
    acc.cell(r, 10, f"=I{r}/$F$6").number_format = "0.0%"
    acc.cell(r, 11, f"=SUM($J$6:J{r})").number_format = "0.0%"
acc["H16"] = "Total"
acc["I16"] = "=SUM(I6:I15)"
acc["J16"] = "=SUM(J6:J15)"
acc["J16"].number_format = "0.0%"
chart(acc, 8, 9, 6, 15, "H19", "Accident locations (10 km bins)", 13.6, 8.4)
note(acc, "B37:N37", "Source: Pem-Sim Kelas E.pdf, p. 1 (embedded data image). Sample statistics use n - 1; quartiles use Excel QUARTILE inclusive method.")
acc.row_dimensions[37].height = 29

# Question 1: fully specified uniform distribution on the stated road length.
base_sheet(ks, "1  Kolmogorov-Smirnov test", "H0: accident location follows Uniform(0, 100 km) | two-sided test, alpha = 0.05", 12, 39)
for col, width in {"B": 7, "C": 12, "D": 12, "E": 13, "F": 13, "G": 13,
                   "H": 3, "I": 18, "J": 17, "K": 12, "L": 12}.items():
    ks.column_dimensions[col].width = width
header(ks, 6, 2, ["Rank i", "Sorted x (km)", "F0(x) = x/100", "(i-1)/n", "i/n", "Local D"])
body(ks, 7, 36, 2, 7)
for r in range(7, 37):
    ks.cell(r, 2, r - 6)
    ks.cell(r, 3, f"=SMALL('1 Accident Data'!$C$6:$C$35,B{r})").number_format = "0.0"
    ks.cell(r, 4, f"=C{r}/100").number_format = "0.0000"
    ks.cell(r, 5, f"=(B{r}-1)/$J$7").number_format = "0.0000"
    ks.cell(r, 6, f"=B{r}/$J$7").number_format = "0.0000"
    ks.cell(r, 7, f"=MAX(D{r}-E{r},F{r}-D{r})").number_format = "0.0000"
header(ks, 6, 9, ["Test result", "Value"])
body(ks, 7, 12, 9, 10)
for r, label, formula, fmt in [
    (7, "n", "='1 Accident Data'!F6", "0"),
    (8, "Alpha", .05, "0.00"),
    (9, "D statistic", "=MAX(G7:G36)", "0.000000"),
    (10, "Exact D critical", .2417034706, "0.000000"),
    (11, "Decision", '=IF(J9>J10,"Reject H0","Do not reject H0")', "General"),
    (12, "Conclusion", '=IF(J9>J10,"Non-uniform evidence","Uniform fit plausible")', "General"),
]:
    ks.cell(r, 9, label)
    ks.cell(r, 10, formula).number_format = fmt
ks["J10"].comment = Comment(
    "Exact two-sided one-sample KS 95th-percentile critical value for n=30: "
    "scipy.stats.kstwo.ppf(0.95, 30) = 0.24170347059707345. "
    "The Uniform(0,100) parameters are specified by the assignment, not estimated from data.",
    "Method",
)
note(ks, "I15:L18", "D is the maximum vertical difference between empirical and Uniform(0,100) CDFs. Reject H0 only when D exceeds the exact 5% critical value.")
note(ks, "B38:L38", "Data: Pem-Sim Kelas E.pdf, p. 1. Method: Law, Simulation Modeling and Analysis, 5th ed., ch. 6. Exact critical value: SciPy kstwo, n = 30.")
ks.row_dimensions[38].height = 29

# Question 2: source observations, summary, and equal-width histogram.
base_sheet(time, "2  Employee recording times", "50 employee observations | time in minutes", 17, 35)
for col, width in {"B": 9, "C": 13, "D": 3, "E": 9, "F": 13, "G": 3,
                   "H": 19, "I": 13, "J": 3, "K": 10, "L": 10, "M": 10,
                   "N": 10, "O": 10, "P": 10, "Q": 8, "T": 13}.items():
    time.column_dimensions[col].width = width
time.column_dimensions["T"].hidden = True
header(time, 5, 2, ["Employee", "Time (min)"])
header(time, 5, 5, ["Employee", "Time (min)"])
body(time, 6, 30, 2, 3)
body(time, 6, 30, 5, 6)
for idx, value in enumerate(TIMES, 1):
    r = (idx - 1) % 25 + 6
    no_col, data_col = (2, 3) if idx <= 25 else (5, 6)
    time.cell(r, no_col, idx)
    c = time.cell(r, data_col, value)
    c.number_format = "0.00"
    c.font = Font(name="Arial", size=9, color=INPUT)
    time.cell(idx + 5, 20, f"={get_column_letter(data_col)}{r}")

header(time, 5, 8, ["Statistical summary", "Value"])
summary_t = [
    ("Count", "=COUNT(T6:T55)", "0"),
    ("Sum (min)", "=SUM(T6:T55)", "0.00"),
    ("Mean (min)", "=AVERAGE(T6:T55)", "0.000"),
    ("Median (min)", "=MEDIAN(T6:T55)", "0.000"),
    ("One mode (min)", "=MODE(T6:T55)", "0.00"),
    ("Minimum (min)", "=MIN(T6:T55)", "0.00"),
    ("Maximum (min)", "=MAX(T6:T55)", "0.00"),
    ("Range (min)", "=I12-I11", "0.00"),
    ("Q1 (min)", "=QUARTILE(T6:T55,1)", "0.000"),
    ("Q3 (min)", "=QUARTILE(T6:T55,3)", "0.000"),
    ("IQR (min)", "=I15-I14", "0.000"),
    ("Sample variance", "=VAR(T6:T55)", "0.000"),
    ("Sample SD (min)", "=STDEV(T6:T55)", "0.000"),
    ("CV", "=I18/I8", "0.0%"),
    ("Skewness", "=SKEW(T6:T55)", "0.000"),
    ("Excess kurtosis", "=KURT(T6:T55)", "0.000"),
]
body(time, 6, 21, 8, 9)
for row, (label, formula, fmt) in enumerate(summary_t, 6):
    time.cell(row, 8, label)
    time.cell(row, 9, formula).number_format = fmt
header(time, 23, 8, ["Interval (min)", "Frequency"])
body(time, 24, 29, 8, 9)
for j in range(6):
    r = j + 24
    time.cell(r, 8, f"{j}-{j+1}" if j < 5 else "5 and above")
    time.cell(r, 9, f'=COUNTIFS($T$6:$T$55,">="&{j},$T$6:$T$55,"<"&{j+1})' if j < 5
              else '=COUNTIF($T$6:$T$55,">=5")')
chart(time, 8, 9, 24, 29, "K6", "Recording times (1 min bins)", 12.2, 10.5)
note(time, "B32:Q32", "Source: Pem-Sim Kelas E.pdf, pp. 1-2 (embedded data image). Sample statistics use n - 1; quartiles use Excel QUARTILE inclusive method.")
note(time, "B33:Q33", "Several values tie for mode: 0.26, 0.36, 0.53, 0.80, and 2.03 min. The MODE cell displays one tied value.")
time.row_dimensions[32].height = 26
time.row_dimensions[33].height = 26

# Question 2: six equiprobable exponential intervals.
base_sheet(chi, "2  Chi-square goodness-of-fit", "H0: recording time follows an exponential distribution | k = 6 | alpha = 0.05", 12, 26)
for col, width in {"B": 9, "C": 17, "D": 17, "E": 14, "F": 14, "G": 16,
                   "H": 3, "I": 21, "J": 20, "K": 11, "L": 9}.items():
    chi.column_dimensions[col].width = width
header(chi, 6, 2, ["Class", "Lower (min)", "Upper (min)", "Observed O", "Expected E", "(O-E)^2/E"])
body(chi, 7, 12, 2, 7)
for r in range(7, 13):
    cls = r - 6
    chi.cell(r, 2, cls)
    chi.cell(r, 3, 0 if r == 7 else f"=D{r-1}").number_format = "0.0000"
    chi.cell(r, 4, f"=-$J$9*LN(1-B{r}/$J$7)" if r < 12 else "infinity")
    if r < 12:
        chi.cell(r, 4).number_format = "0.0000"
        chi.cell(r, 5, f'=COUNTIFS(\'2 Time Data\'!$T$6:$T$55,">="&C{r},\'2 Time Data\'!$T$6:$T$55,"<"&D{r})')
    else:
        chi.cell(r, 5, '=COUNTIF(\'2 Time Data\'!$T$6:$T$55,">="&C12)')
    chi.cell(r, 6, "=$J$8/$J$7").number_format = "0.0000"
    chi.cell(r, 7, f"=(E{r}-F{r})^2/F{r}").number_format = "0.0000"
chi["D12"].alignment = Alignment(horizontal="left")
for col, value in [(2, "Total"), (5, "=SUM(E7:E12)"), (6, "=SUM(F7:F12)"), (7, "=SUM(G7:G12)")]:
    chi.cell(13, col, value)
    chi.cell(13, col).font = Font(name="Arial", size=9, bold=True, color=NAVY)
    chi.cell(13, col).fill = PatternFill("solid", fgColor=PALE)
chi["F13"].number_format = "0.0"
chi["G13"].number_format = "0.0000"

header(chi, 6, 9, ["Test parameter", "Value"])
body(chi, 7, 15, 9, 10)
tests = [
    (7, "Classes k", 6, "0"),
    (8, "Sample size n", "='2 Time Data'!I6", "0"),
    (9, "Estimated mean", "='2 Time Data'!I8", "0.0000"),
    (10, "Rate lambda = 1/mean", "=1/J9", "0.0000"),
    (11, "Degrees of freedom", "=J7-1-1", "0"),
    (12, "Chi-square statistic", "=G13", "0.0000"),
    (13, "Critical value (5%)", "=CHIINV(0.05,J11)", "0.0000"),
    (14, "p-value", "=CHIDIST(J12,J11)", "0.0000"),
    (15, "Decision", '=IF(J12>J13,"Reject H0","Do not reject H0")', "General"),
]
for r, label, formula, fmt in tests:
    chi.cell(r, 9, label)
    chi.cell(r, 10, formula).number_format = fmt
note(chi, "B17:L19", "Class boundaries are fitted exponential quantiles: x_i = -mean * LN(1 - i/6), i = 1,...,5. Each class has expected frequency 50/6 = 8.33, satisfying the usual chi-square minimum of 5.")
note(chi, "B20:L22", "One parameter (the mean) is estimated from the same sample, so df = 6 - 1 - 1 = 4. A non-rejection means the exponential model is plausible for these data; it does not prove the model is true.")
note(chi, "B25:L25", "Data: Pem-Sim Kelas E.pdf, p. 2. Method: Law, Simulation Modeling and Analysis, 5th ed., ch. 6; F(x) = 1 - EXP(-x/mean).")
chi.row_dimensions[25].height = 27

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(OUTPUT)
print(OUTPUT)
